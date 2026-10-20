from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE=Path(__file__).parent
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=Inches(.65)
s.left_margin=s.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(11 if name=='Normal' else 24 if name=='Title' else 19 if name=='Heading 1' else 13 if name=='Heading 2' else 9)
 st.paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('COMASA  |  Instructivo de uso  |  ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
for r in footer.runs:r.font.size=Pt(9)
def p(text):
 para=doc.add_paragraph()
 for i,part in enumerate(text.split('**')): para.add_run(part).bold=bool(i%2)
 return para
def shot(name,box=None,width=7.1):
 path=BASE/'instructivo_capturas'/name; w,h=Image.open(path).size
 x,y,r,b=box or (0,0,w,h)
 para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
 shape=para.add_run().add_picture(str(path),width=Inches(width),height=Inches(width*(b-y)/(r-x)))
 if box:
  rect=OxmlElement('a:srcRect')
  for k,v in {'l':x/w,'t':y/h,'r':(w-r)/w,'b':(h-b)/h}.items():rect.set(k,str(round(v*100000)))
  shape._inline.graphic.graphicData.pic.blipFill.insert(1,rect)
 shape._inline.docPr.set('descr',name.replace('.png','').replace('-',' '))
 para.paragraph_format.space_after=Pt(5)
def page(title):doc.add_page_break(); doc.add_heading(title,1)

doc.add_heading('Instructivo del Cotizador COMASA',0)
p('Guía para asesores de ventas y personal administrativo. Siga el flujo: seleccionar cliente, agregar productos, revisar totales y emitir la cotización. Edición del 5 de octubre de 2026. Las capturas originales son ilustrativas; las condiciones vigentes de kits se explican en la sección 10.')
shot('02-inicio.png',width=6.4)
doc.add_heading('Acceso y navegación',2)
p('Ingrese su **Correo** y **Clave** y pulse **Ingresar**. Para recuperar el acceso, escriba su correo y use **Restablecer clave**; la nueva clave requiere al menos ocho caracteres.')
p('Los asesores acceden a **Cotización**. El perfil administrativo también dispone de **Inicio**, **Administración** y **Gestión de cotizaciones**. Use **Cerrar sesión** al terminar.')
p('**Contenido:** cliente y productos, importación, revisión y emisión, consulta, administración, configuración de ofertas, kits por SET y ayuda.')

page('1 Seleccionar el cliente')
shot('04-clientes.png',(115,145,870,662))
p('1. En **Cotización**, pulse **Buscar cliente**.\n2. Escriba nombre, cédula o celular.\n3. Seleccione el resultado correcto.\n4. Revise los datos y el segmento base en los parámetros de cotización.')
p('**Segmento base:** se obtiene del cliente seleccionado y se utiliza para calcular sus condiciones comerciales. Seleccione el cliente antes de consultar ofertas.')
p('**Limpiar** quita el cliente y su segmento, pero conserva los productos. Para cambiar de cliente, vuelva a buscarlo y revise los precios recalculados.')
p('La captura muestra una búsqueda ilustrativa sin resultados. Si esto ocurre, pruebe otro dato del cliente o solicite revisar su registro en la base.')

page('2 Buscar y agregar productos')
shot('05-productos.png',(15,57,970,750))
p('1. Pulse **Agregar SKU**.\n2. Use **División** o busque por código, descripción o número de parte.\n3. Seleccione un producto y escriba la cantidad en **Cant.**\n4. Revise precio, existencia y condiciones de las ofertas.\n5. Pulse **Agregar SKU**, **Agregar SKU con oferta** o **Agregar SKU (Precio de lista)**, según se muestre.')
p('Repita para otros productos y cierre el buscador. Sin cliente seleccionado, se presenta el precio de lista. Las existencias corresponden a la tienda 1041 cuando el inventario está activo.')
p('En kits, revise los SET y sus cantidades. Cada SET admite los SKU elegibles indicados. Todos deben completarse con unidades diferentes. Los kits y recompensas se calculan con los productos agregados a la cotización.')

page('3 Importar SKU desde archivo o texto')
shot('06-importar.png',(210,232,777,575),width=5.2)
p('**Archivo:** pulse **Importar → Archivo → Seleccionar archivo**. Se admiten XLSX, CSV y TSV. Un archivo válido **reemplaza las líneas actuales** de la cotización.')
shot('07-pegar.png',(210,180,777,625),width=5.0)
p('**Pegar datos:** pegue SKU y cantidad, una línea por producto. Revise el contador y pulse **Agregar**. Esta opción **añade líneas** a las existentes. Ejemplo: código 100620821, cantidad 2.')

page('4 Revisar productos y totales')
shot('08-revision.png',(275,175,1475,925))
p('1. Verifique cliente, segmento y productos.\n2. Ajuste **Cantidad** directamente en la tabla. Use la papelera para eliminar una línea.\n3. Espere a que termine el cálculo de promociones.\n4. Revise **Subtotal final**, **IVA**, **Total ahorrado** y **Total con IVA**.')
p('**Ejemplo de la captura:** 2 unidades a C$451.30 suman C$902.60 antes de IVA; el total mostrado es C$1,037.99. Estos importes ilustran la pantalla y pueden variar con el catálogo y las promociones vigentes.')
p('Los avisos debajo de los botones indican qué falta para emitir. Si las promociones siguen cargando o hay un error de cálculo, resuelva el aviso antes de continuar.')

page('5 Descargar borrador o emitir')
shot('09-emision.png',(220,225,762,580))
p('**Borrador PDF:** descarga un documento de revisión sin asignar consecutivo. Requiere productos y un cálculo disponible.')
p('**Emitir cotización:** requiere un cliente y al menos un SKU. Pulse el botón, revise cliente, segmento, líneas y total; después pulse **Confirmar emisión**. La aplicación registra la cotización, asigna consecutivo y descarga el PDF oficial.')
p('Use **Cancelar** para volver a revisar. La captura corresponde al paso previo a confirmar.')
p('Si la app informa que la cotización se emitió pero el PDF no se pudo descargar, anote el consecutivo y consulte al administrador antes de emitir nuevamente.')

page('6 Consultar cotizaciones emitidas')
shot('12-gestion.png',(270,75,1475,885))
p('**Solo perfil administrativo.**\n1. Abra **Gestión de cotizaciones**.\n2. Escriba un ID, cliente, usuario o segmento y ajuste **Buscar en**.\n3. Delimite las fechas con **Desde** y **Hasta** y pulse **Buscar**.\n4. Seleccione una fila para revisar el detalle de sus productos y precios.')
p('**Exportar a Excel** descarga las cotizaciones consultadas. En el detalle, **Excel** descarga una cotización individual con sus códigos, cantidades y precios. **Limpiar** restablece los filtros.')
p('La captura ilustra una consulta sin coincidencias; por eso la exportación está deshabilitada. Los borradores sin emisión no generan consecutivo.')

page('7 Cargar y actualizar datos')
shot('10-administracion.png',(270,55,1475,815),width=6.6)
p('**Solo perfil administrativo.** Seleccione el proceso, cargue el archivo completo, revise el conteo y la vista previa y pulse su botón de publicación. Espere el mensaje de resultado.')
table=doc.add_table(rows=1, cols=2); table.style='Light Shading Accent 1'
table.rows[0].cells[0].text='Proceso';table.rows[0].cells[1].text='Botón de publicación'
for a,b in [('Promociones','Sincronizar'),('Clientes','Actualizar clientes'),('Catálogo','Actualizar catálogo'),('Inventario','Actualizar inventario'),('Tiendas','Actualizar tiendas')]:
 c=table.add_row().cells;c[0].text=a;c[1].text=b
p('**Clientes, catálogo, inventario y tiendas reemplazan toda la base del proceso.** Seleccionar el archivo permite revisarlo; publicar requiere pulsar actualizar. Inventario conserva registros de tienda 1041. En promociones revise el modo **Carga total** o **Actualizacion** antes de sincronizar.')

page('8 Configurar combinación y cantidades')
shot('11-ofertas.png',(520,340,1455,1100))
p('1. En **Administración**, abra **Config. ofertas**.\n2. Filtre por **ID promo**, **ID oferta** o **SKU** y pulse **Buscar**.\n3. Identifique la fila correcta por oferta, código y segmento.')
p('**Combina:** define si admite combinación. La casilla se guarda inmediatamente y afecta las filas de la promoción y oferta correspondientes.')
p('**Threshold por fila:** se usa en ofertas que no son kits. Algunas modalidades fijan el tipo mínimo y permiten cero. Pulse **Guardar** en la fila. En kits, los umbrales se editan dentro de cada SET mediante **Definir regla** o **Editar regla**.')
p('Desplace horizontalmente la tabla si no ve las últimas columnas o el botón **Guardar**. Use **Definir regla** o **Editar regla** para abrir las modalidades de la siguiente página.')

page('9 Definir una regla comercial')
shot('13-regla.png',(420,430,1380,920))
p('Seleccione **Modalidad**, complete los campos de la condición comercial y pulse **Guardar regla**. **Cerrar** permite salir sin guardar.')
for text in ['**Unitario importado:** utiliza el precio o descuento del reporte.','**Paquete de un SKU:** indique unidades por paquete y precio TOTAL del paquete.','**Kit por SET:** agrupe los SKU elegibles, indique unidades exactas o mínimas y el beneficio de cada SET. Vea el ejemplo de la sección 10.','**Unidades mezcladas:** indique los SKU elegibles, unidades por grupo y beneficio.','**Compra X y beneficio en Y:** complete códigos y cantidades de compra y recompensa. El producto de recompensa debe estar en la cotización.']:
 p(text)
p('En porcentajes, **100 significa gratis**. Las unidades consumidas en un paquete no califican nuevamente en otro paquete. Revise cualquier mensaje de validación antes de guardar.')

page('10 Configurar kits por SET')
p('**Configuración:** abra la oferta y use **Agregar SET**. Indique sus SKU separados por coma, unidades requeridas y tipo de umbral. El beneficio puede ser **Sin beneficio**, **Porcentaje** (100 = gratis) o **Precio unitario (C$)**. Complete todos los SET y pulse **Guardar regla**.')
p('**Ejemplo de oferta 23456:** SET 1 acepta A, B o C y exige 3 unidades exactas, sin beneficio propio. SET 2 exige 2 unidades exactas de D y otorga 100 % de descuento. Son válidos 3 A + 2 D, 1 A + 2 C + 2 D o 1 A + 1 B + 1 C + 2 D. Con una sola D, el kit no se completa.')
p('**Umbral exacto:** si pide 18 A y se agregan 20, puede asignar 18 al kit. Las 2 restantes buscan otra oferta válida para esas 2 unidades; si no existe, se calculan a precio regular.')
p('**Umbral mínimo:** si exige 30 unidades y concede 10 %, con 50 unidades pueden beneficiarse las 50. Con 29 no se cumple. También deben completarse los otros SET. El motor compara el resultado con otras ofertas para elegir el menor total.')
p('**SKU compartidos:** un SKU puede pertenecer a varios SET, pero una unidad no puede contarse dos veces. Si un SET exige 3 A y otro 2 A, se necesitan 5 A. Cada unidad recibe el beneficio del SET al que se asigne.')
p('**Combina:** se controla para la promoción y oferta completas desde la tabla. Conserva la combinación existente de una base no combinable con beneficios combinables; un kit no combinable puede recibir un descuento unitario marcado como combinable.')
p('**Pendientes:** se conservan kits de cualquier número de SKU. Sin SET configurados, no se aplican automáticamente. Las reglas antiguas deben revisarse y guardarse como SET. El reporte no identifica a qué SET pertenece cada código.')
p('**Revisión:** las cantidades por SET son enteras positivas. Las regalías deben estar agregadas. El desglose identifica las unidades por SET; completar un kit no obliga a aplicarlo si otra combinación resulta más económica.')

page('11 Resolver avisos y terminar')
shot('01-acceso.png',(15,140,415,398),width=4.4)
for text in ['**Perfil pendiente:** solicite al administrador revisar la autorización de su usuario.','**Cliente sin resultados:** pruebe nombre, cédula o celular y confirme que el registro exista.','**Archivo sin SKU válidos:** revise formato, encabezados y códigos; seleccione el archivo corregido.','**Promociones cargando o error de cálculo:** espere la carga; si persiste, revise la conexión y solicite comprobar los datos.','**Emitir deshabilitado:** lea el aviso de requisitos debajo de los botones.','**Actualización fallida:** revise el mensaje y no dé por publicada la información hasta obtener confirmación de éxito.']:
 p(text)
doc.add_heading('Comprobación antes de emitir',2)
p('Cliente correcto · Productos correctos · Cantidades revisadas · Ofertas calculadas · Total con IVA verificado.')
p('Al finalizar su trabajo, use **Cerrar sesión**. Conserve el PDF oficial y su consecutivo para identificar la cotización.')
doc.core_properties.title='Instructivo del Cotizador COMASA'
doc.core_properties.subject='Guía de operación actualizada con kits por SET'
doc.core_properties.author='COMASA'
for element in list(doc.styles.element.iter(qn('w:pBdr'))) + list(doc.element.iter(qn('w:pBdr'))):
 element.getparent().remove(element)
doc.save(BASE/'Instructivo_Cotizador_COMASA.docx')
print(BASE/'Instructivo_Cotizador_COMASA.docx')
